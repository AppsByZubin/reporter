from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from common.constants import DEFAULT_TEMPLATE
from utils.excel_utils import write_report


class ExcelUtilsTests(TestCase):
    def test_report_removes_sections_for_bots_without_orders(self) -> None:
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        report_data = {
            "firebot": ([{"trade_id": "fire-order"}], ""),
            "trendobot": ([], ""),
            "titanbot": ([{"trade_id": "titan-order"}], ""),
        }
        with TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "production.xlsx"
            write_report(
                DEFAULT_TEMPLATE,
                output_path,
                ["firebot", "trendobot", "titanbot"],
                report_data,
            )
            workbook = load_workbook(output_path)

        worksheet = workbook.active
        bot_titles = {
            str(cell.value).strip().lower()
            for cell in worksheet["E"]
            if cell.value
            and str(cell.value).strip().lower().endswith("bot")
        }
        self.assertEqual(bot_titles, {"firebot", "titanbot"})
        self.assertEqual(set(worksheet.tables), {"Table26", "Table274"})

    def test_report_writes_meanbot_after_removing_preceding_sections(self) -> None:
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        report_data = {
            "meanbot": ([{"trade_id": "mean-order", "amount": "125.50"}], ""),
        }
        with TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "production.xlsx"
            write_report(
                DEFAULT_TEMPLATE,
                output_path,
                ["meanbot"],
                report_data,
            )
            workbook = load_workbook(output_path)

        worksheet = workbook.active
        self.assertEqual(worksheet["E3"].value, "Meanbot")
        self.assertEqual(worksheet["E6"].value, "mean-order")
        self.assertEqual(worksheet["P13"].value, 125.5)
        self.assertEqual(set(worksheet.tables), {"Table27428"})
